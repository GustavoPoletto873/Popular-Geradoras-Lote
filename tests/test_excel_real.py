"""Excel REAL via COM (Windows com Office). Só roda com `--run-excel`:

    python -m pytest tests/test_excel_real.py --run-excel -q

Usa uma boleta-modelo SINTÉTICA (.xlsx) e os arquivos de exemplo de tests/dourados como insumos, com os nomes
canônicos que o pipeline produz. Não toca nenhuma pasta de produção (tudo em tmp_path).
"""

import shutil
from pathlib import Path

import openpyxl
import pytest

from contabilidade_mensal.integrations.excel.adapter import PopuladorExcel
from contabilidade_mensal.integrations.excel.origem import OrigemLocal
from populador import config as cfg_pop

pytestmark = pytest.mark.excel

DOURADOS = Path(__file__).parent / "dourados"
CNPJ = "12345678000190"
ABAS = [*cfg_pop.ABAS_PARA_LIMPAR, "Carteira inicial", cfg_pop.ABA_MAPA_CONTABIL]


def valores(caminho: Path, aba: str, faixa: str):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    try:
        return [[c.value for c in linha] for linha in wb[aba][faixa]]
    finally:
        wb.close()


@pytest.fixture
def cenario(tmp_path, criar_fundo):
    fundo = criar_fundo(1001, nome="FII KRONOS", tipo="FII", exercicio_mes=11)

    modelo = openpyxl.Workbook()
    modelo.remove(modelo.active)
    for aba in ABAS:
        modelo.create_sheet(aba)
    modelo["Carteira final"]["B2"] = "VALOR-DO-MES-ANTERIOR"
    modelo["Extrato"]["A1"] = "EXTRATO-VELHO"
    pasta_anterior = tmp_path / "G" / "FII" / "FII KRONOS" / "Data Base Novembro 2026" / "202607"
    pasta_anterior.mkdir(parents=True)
    modelo.save(pasta_anterior / "FII KRONOS 202607.xlsx")

    balancete = openpyxl.Workbook()
    balancete.active["B1"] = "BALANCETE-NOVO"
    balancete.active["B2"] = 123.45
    insumos = tmp_path / "insumos"
    insumos.mkdir()
    balancete.save(insumos / f"202608_{CNPJ}_BalanceteContabilFinal.xlsx")
    for cenario, nome in [
        ("carteira_final", f"202608_{CNPJ}_CarteiraFinal.xlsx"),
        ("carteira_inicial", f"202608_{CNPJ}_CarteiraInicial.xlsx"),
        ("extrato_padrao", f"202608_{CNPJ}_ExtratoCC.xlsx"),
        ("mov_uma_classe", f"202608_{CNPJ}_MovCotista.xlsx"),
        ("posicao_final_data_col4", f"202608_{CNPJ}_SaldoAplicacaoCotistaFinal.xlsx"),
    ]:
        origem = next((DOURADOS / cenario / "saida").glob("*.xlsx"))
        shutil.copy2(origem, insumos / nome)
    return fundo, sorted(insumos.iterdir()), tmp_path


def test_popula_a_boleta_com_os_insumos_canonicos_e_limpa_o_modelo(cenario, competencia):
    fundo, arquivos, tmp_path = cenario
    gerada = PopuladorExcel(OrigemLocal(tmp_path / "G")).popular(
        fundo=fundo, competencia=competencia, arquivos=arquivos, pasta_trabalho=tmp_path / "trab"
    )

    assert gerada.caminho.name == "FII KRONOS 202608.xlsx"  # extensão herdada da boleta anterior
    assert gerada.caminho.is_file()
    assert not any("não encontrado" in a.lower() for a in gerada.avisos), gerada.avisos

    # o que era do mês anterior sumiu e o novo entrou (comparado com o arquivo de origem, célula a célula)
    final = next((DOURADOS / "carteira_final" / "saida").glob("*.xlsx"))
    esperado = valores(final, openpyxl.load_workbook(final).sheetnames[0], "B1:H6")
    assert valores(gerada.caminho, "Carteira final", "B1:H6") == esperado
    assert valores(gerada.caminho, "Carteira final", "B2:B2") != [["VALOR-DO-MES-ANTERIOR"]]
    assert valores(gerada.caminho, "Extrato", "A1:A1") != [["EXTRATO-VELHO"]]
    assert valores(gerada.caminho, "Balancete", "B1:B2") == [["BALANCETE-NOVO"], [123.45]]


def test_o_modelo_do_mes_anterior_nao_e_alterado(cenario, competencia):
    fundo, arquivos, tmp_path = cenario
    modelo = tmp_path / "G" / "FII" / "FII KRONOS" / "Data Base Novembro 2026" / "202607" / "FII KRONOS 202607.xlsx"
    antes = modelo.read_bytes()
    PopuladorExcel(OrigemLocal(tmp_path / "G")).popular(
        fundo=fundo, competencia=competencia, arquivos=arquivos, pasta_trabalho=tmp_path / "trab"
    )
    assert modelo.read_bytes() == antes


def test_boleta_xlsb_herda_a_extensao_e_e_gravada_pelo_proprio_excel(cenario, competencia):
    """O formato real das boletas é .xlsb (openpyxl não escreve): confere a ida e volta pelo Excel."""
    from populador.excel_app import ExcelSession

    fundo, arquivos, tmp_path = cenario
    pasta = tmp_path / "G" / "FII" / "FII KRONOS" / "Data Base Novembro 2026" / "202607"
    xlsx = pasta / "FII KRONOS 202607.xlsx"
    XL_XLSB = 50
    with ExcelSession() as sessao:
        wb = sessao.open_workbook(xlsx, read_only=False)
        wb.SaveAs(str(pasta / "FII KRONOS 202607.xlsb"), XL_XLSB)
        wb.Close(SaveChanges=False)
    xlsx.unlink()

    gerada = PopuladorExcel(OrigemLocal(tmp_path / "G")).popular(
        fundo=fundo, competencia=competencia, arquivos=arquivos, pasta_trabalho=tmp_path / "trab"
    )
    assert gerada.caminho.name == "FII KRONOS 202608.xlsb"
    with ExcelSession() as sessao:
        wb = sessao.open_workbook(gerada.caminho, read_only=True)
        try:
            assert wb.Sheets("Balancete").Range("B1").Value == "BALANCETE-NOVO"
            assert wb.Sheets("Carteira final").Range("B2").Value == 111
        finally:
            wb.Close(SaveChanges=False)
